# mypy: ignore-errors
from __future__ import annotations

from datetime import datetime
from pathlib import Path

from adv_archon.core.document_draft import DocumentDraft, DraftSection, DraftSource, DraftTable


def generate_expediente_docx(
    draft: DocumentDraft,
    output_path: Path,
    *,
    template_path: Path | None = None,
    client_logo_path: Path | None = None,
    archon_logo_path: Path | None = None,
) -> Path:
    """Generate a real DOCX report from an editable DocumentDraft."""

    return _generate_docx(
        draft=draft,
        output_path=output_path,
        template_path=template_path,
        client_logo_path=client_logo_path,
        archon_logo_path=archon_logo_path,
        heading="Informe preliminar de viabilidad urbanistica",
    )


def generate_research_docx(
    draft: DocumentDraft,
    output_path: Path,
    *,
    template_path: Path | None = None,
) -> Path:
    return _generate_docx(
        draft=draft,
        output_path=output_path,
        template_path=template_path,
        heading="Informe de investigacion",
    )


def generate_budget_docx(
    draft: DocumentDraft,
    output_path: Path,
    *,
    template_path: Path | None = None,
) -> Path:
    return _generate_docx(
        draft=draft,
        output_path=output_path,
        template_path=template_path,
        heading="Presupuesto y alcance profesional",
    )


def _generate_docx(
    *,
    draft: DocumentDraft,
    output_path: Path,
    heading: str,
    template_path: Path | None = None,
    client_logo_path: Path | None = None,
    archon_logo_path: Path | None = None,
) -> Path:
    from docx import Document  # type: ignore[import-untyped]
    from docx.enum.text import WD_ALIGN_PARAGRAPH  # type: ignore[import-untyped]
    from docx.shared import Inches, Pt, RGBColor  # type: ignore[import-untyped]

    output_path = output_path.expanduser()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document = Document(str(template_path)) if template_path else Document()
    _configure_styles(document)

    header = document.add_paragraph()
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if archon_logo_path and archon_logo_path.exists():
        header.add_run().add_picture(str(archon_logo_path), width=Inches(0.75))
    if client_logo_path and client_logo_path.exists():
        header.add_run("   ")
        header.add_run().add_picture(str(client_logo_path), width=Inches(0.75))

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("ADV ARCHON")
    title_run.bold = True
    title_run.font.size = Pt(18)
    title_run.font.color.rgb = RGBColor(201, 162, 39)
    subtitle = document.add_paragraph(heading)
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

    document.add_heading(draft.title, level=1)
    _add_key_value_table(
        document,
        (
            ("Cliente", draft.client_name),
            ("Proyecto", draft.project_name),
            ("Municipio", draft.municipality),
            ("Referencia catastral", draft.cadastral_ref),
            ("Veredicto preliminar", draft.verdict),
            ("Fecha", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ),
    )
    _add_notice(document)
    document.add_heading("Resumen ejecutivo", level=1)
    document.add_paragraph(draft.executive_summary or "Pendiente de completar.")

    for section in draft.sections:
        if section.include_in_report:
            _add_section(document, section)
    for table in draft.tables:
        if table.include_in_report:
            _add_table(document, table)
    if draft.sources:
        document.add_heading("Fuentes y trazabilidad", level=1)
        _add_sources_table(document, draft.sources)
    if draft.warnings:
        document.add_heading("Advertencias", level=1)
        for warning in draft.warnings:
            document.add_paragraph(warning, style="List Bullet")
    if draft.next_steps:
        document.add_heading("Proximos pasos", level=1)
        for step in draft.next_steps:
            document.add_paragraph(step, style="List Number")

    document.save(str(output_path))
    return output_path


def _configure_styles(document: object) -> None:
    styles = document.styles  # type: ignore[attr-defined]
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = 10
    for style_name in ("Title", "Heading 1", "Heading 2"):
        style = styles[style_name]
        style.font.name = "Georgia"
        style.font.color.rgb = _rgb(46, 46, 44)


def _add_notice(document: object) -> None:
    paragraph = document.add_paragraph()  # type: ignore[attr-defined]
    run = paragraph.add_run(
        "Documento preliminar. Las conclusiones deben ser revisadas por tecnico "
        "competente antes de su uso profesional o juridico."
    )
    run.italic = True


def _add_section(document: object, section: DraftSection) -> None:
    document.add_heading(section.title, level=min(max(section.level, 1), 3))  # type: ignore[attr-defined]
    for block in section.content.split("\n\n"):
        if block.strip():
            document.add_paragraph(block.strip())  # type: ignore[attr-defined]


def _add_key_value_table(document: object, rows: tuple[tuple[str, str], ...]) -> None:
    table = document.add_table(rows=1, cols=2)  # type: ignore[attr-defined]
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "Campo"
    table.rows[0].cells[1].text = "Valor"
    for key, value in rows:
        cells = table.add_row().cells
        cells[0].text = key
        cells[1].text = value or "-"


def _add_table(document: object, draft_table: DraftTable) -> None:
    document.add_heading(draft_table.title, level=2)  # type: ignore[attr-defined]
    column_count = max(1, len(draft_table.columns))
    table = document.add_table(rows=1, cols=column_count)  # type: ignore[attr-defined]
    table.style = "Table Grid"
    for idx, column in enumerate(draft_table.columns or ("Dato",)):
        table.rows[0].cells[idx].text = column
    for row in draft_table.rows:
        cells = table.add_row().cells
        for idx in range(column_count):
            cells[idx].text = row[idx] if idx < len(row) else ""


def _add_sources_table(document: object, sources: tuple[DraftSource, ...]) -> None:
    table = document.add_table(rows=1, cols=5)  # type: ignore[attr-defined]
    table.style = "Table Grid"
    headers = ("Fuente", "Tipo", "Referencia", "Resultado", "Confianza")
    for idx, header in enumerate(headers):
        table.rows[0].cells[idx].text = header
    for source in sources:
        cells = table.add_row().cells
        cells[0].text = source.title
        cells[1].text = "Oficial" if source.official else source.source_type
        cells[2].text = source.url or source.reference
        cells[3].text = source.result
        cells[4].text = source.confidence


def _rgb(red: int, green: int, blue: int) -> object:
    from docx.shared import RGBColor  # type: ignore[import-untyped]

    return RGBColor(red, green, blue)
