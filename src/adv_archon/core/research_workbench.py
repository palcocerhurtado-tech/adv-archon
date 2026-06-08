from __future__ import annotations

import re
from collections.abc import Callable, Sequence
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

from docx import Document
from fpdf import FPDF, XPos, YPos

from adv_archon.tools.files import read_file
from adv_archon.tools.web import web_fetch, web_search

SearchFn = Callable[[str, int], Any]
FetchFn = Callable[[str], Any]


@dataclass(frozen=True, slots=True)
class ResearchAttachment:
    path: str
    title: str
    content: str

    def as_payload(self) -> dict[str, str]:
        return {"path": self.path, "title": self.title, "content": self.content}


@dataclass(frozen=True, slots=True)
class ResearchEvidence:
    id: str
    title: str
    url: str
    query: str
    excerpt: str
    source_type: str = "web"

    def as_payload(self) -> dict[str, str]:
        return {
            "id": self.id,
            "title": self.title,
            "url": self.url,
            "query": self.query,
            "excerpt": self.excerpt,
            "source_type": self.source_type,
        }


@dataclass(frozen=True, slots=True)
class ResearchWorkbenchResult:
    question: str
    subquestions: tuple[str, ...]
    queries: tuple[str, ...]
    attachments: tuple[ResearchAttachment, ...]
    evidences: tuple[ResearchEvidence, ...]
    formula_candidates: tuple[str, ...]
    synthesis: str
    gaps: tuple[str, ...]
    next_steps: tuple[str, ...]

    @property
    def evidence_count(self) -> int:
        return len(self.evidences)

    def as_payload(self) -> dict[str, Any]:
        return {
            "question": self.question,
            "subquestions": list(self.subquestions),
            "queries": list(self.queries),
            "attachments": [item.as_payload() for item in self.attachments],
            "evidences": [item.as_payload() for item in self.evidences],
            "formula_candidates": list(self.formula_candidates),
            "synthesis": self.synthesis,
            "gaps": list(self.gaps),
            "next_steps": list(self.next_steps),
            "evidence_count": self.evidence_count,
        }


def load_research_attachments(
    paths: Sequence[str],
    *,
    preview: bool = True,
) -> tuple[ResearchAttachment, ...]:
    attachments: list[ResearchAttachment] = []
    for raw_path in paths:
        path = str(raw_path or "").strip()
        if not path:
            continue
        result = read_file(path, preview=preview)
        content = str(result.payload.get("content") or "")
        resolved = str(result.payload.get("path") or path)
        attachments.append(
            ResearchAttachment(
                path=resolved,
                title=Path(resolved).name,
                content=_clip(content, 12000),
            )
        )
    return tuple(attachments)


def build_research_queries(
    question: str,
    *,
    attachments: Sequence[ResearchAttachment] = (),
    max_queries: int = 5,
) -> tuple[str, ...]:
    clean_question = _clean(question)
    terms = _keywords(clean_question)
    attachment_terms: list[str] = []
    for attachment in attachments:
        attachment_terms.extend(_keywords(attachment.content)[:5])
    unique_terms = list(dict.fromkeys([*terms, *attachment_terms]))

    queries = [
        clean_question,
        f"{clean_question} formulas metodo solucion",
        f"{clean_question} ejemplos resueltos",
    ]
    if unique_terms:
        queries.append(" ".join(unique_terms[:8]))
    if any(term in clean_question.casefold() for term in ("arquitect", "urban", "obra")):
        queries.append(f"{clean_question} normativa tecnica espana")
    return tuple(dict.fromkeys(query for query in queries if query.strip()))[:max_queries]


def build_research_subquestions(question: str) -> tuple[str, ...]:
    clean_question = _clean(question)
    return (
        f"Que datos o conceptos exige resolver: {clean_question}?",
        "Que formulas, criterios o metodos verificables aparecen en las fuentes?",
        "Que evidencias sostienen la respuesta y que limites tienen?",
        "Que resultado final se puede entregar y que queda pendiente de validar?",
    )


def run_research_workbench(
    question: str,
    *,
    attachment_paths: Sequence[str] = (),
    search_fn: SearchFn | None = None,
    fetch_fn: FetchFn | None = None,
    search_results_per_query: int = 4,
    fetch_top_results: int = 2,
) -> ResearchWorkbenchResult:
    attachments = load_research_attachments(attachment_paths)
    queries = build_research_queries(question, attachments=attachments)
    subquestions = build_research_subquestions(question)
    search = search_fn or _default_search
    fetch = fetch_fn or _default_fetch

    evidences: list[ResearchEvidence] = []
    seen_urls: set[str] = set()
    for query in queries:
        try:
            results = _extract_search_results(search(query, search_results_per_query))
        except Exception:
            continue
        for item in results[:fetch_top_results]:
            url = str(item.get("url") or "").strip()
            if not url or url in seen_urls:
                continue
            seen_urls.add(url)
            title = str(item.get("title") or url).strip()
            snippet = str(item.get("snippet") or "").strip()
            try:
                fetched = _extract_fetch_text(fetch(url))
            except Exception:
                fetched = ""
            excerpt = _best_excerpt(fetched or snippet, question)
            if not excerpt:
                continue
            evidences.append(
                ResearchEvidence(
                    id=f"E{len(evidences) + 1}",
                    title=_clip(title, 140),
                    url=url,
                    query=query,
                    excerpt=excerpt,
                )
            )

    formula_candidates = tuple(
        dict.fromkeys(
            [
                *extract_formula_candidates(question),
                *[
                    formula
                    for attachment in attachments
                    for formula in extract_formula_candidates(attachment.content)
                ],
                *[
                    formula
                    for evidence in evidences
                    for formula in extract_formula_candidates(evidence.excerpt)
                ],
            ]
        )
    )[:12]
    synthesis = build_research_synthesis(
        question=question,
        attachments=attachments,
        evidences=tuple(evidences),
        formula_candidates=formula_candidates,
    )
    gaps = _build_gaps(evidences=tuple(evidences), attachments=attachments)
    next_steps = _build_next_steps(formula_candidates=formula_candidates, evidences=evidences)
    return ResearchWorkbenchResult(
        question=_clean(question),
        subquestions=subquestions,
        queries=queries,
        attachments=attachments,
        evidences=tuple(evidences),
        formula_candidates=formula_candidates,
        synthesis=synthesis,
        gaps=gaps,
        next_steps=next_steps,
    )


def extract_formula_candidates(text: str) -> tuple[str, ...]:
    candidates: list[str] = []
    formula_pattern = re.compile(
        r"(?i)(?:formula|ecuacion|equation|=|∑|sqrt|raiz|derivada|integral|coeficiente)"
    )
    for line in text.splitlines():
        clean = _clean(line)
        if len(clean) < 6 or len(clean) > 220:
            continue
        if formula_pattern.search(clean):
            candidates.append(clean)
    inline = re.findall(r"[\w\s()^*/+\-.]+=[\w\s()^*/+\-.]+", text)
    candidates.extend(_clean(item) for item in inline if len(_clean(item)) <= 220)
    return tuple(dict.fromkeys(candidates))[:16]


def build_research_synthesis(
    *,
    question: str,
    attachments: Sequence[ResearchAttachment],
    evidences: Sequence[ResearchEvidence],
    formula_candidates: Sequence[str],
) -> str:
    lines = [
        f"Pregunta de trabajo: {_clean(question)}",
        "",
        "Resumen preliminar:",
    ]
    if attachments:
        lines.append(
            f"- Se han considerado {len(attachments)} adjunto(s) aportados por el usuario."
        )
    if evidences:
        lines.append(
            f"- Se han recopilado {len(evidences)} evidencia(s) web para apoyar la respuesta."
        )
        for evidence in evidences[:4]:
            lines.append(f"- {evidence.id}: {_clip(evidence.excerpt, 260)}")
    else:
        lines.append("- No se han recopilado evidencias web; el resultado queda como borrador.")
    if formula_candidates:
        lines.append("")
        lines.append("Formulas o relaciones detectadas:")
        for formula in formula_candidates[:6]:
            lines.append(f"- {formula}")
    lines.append("")
    lines.append(
        "Criterio de salida: redactar respuesta final separando datos encontrados, "
        "formulas/metodo, supuestos y puntos pendientes de verificacion."
    )
    return "\n".join(lines)


def export_research_docx(result: ResearchWorkbenchResult, output_path: Path) -> Path:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.add_heading("Informe de investigacion - ADV ARCHON", level=0)
    doc.add_paragraph(f"Fecha: {datetime.now().strftime('%Y-%m-%d %H:%M')}")
    doc.add_paragraph(f"Pregunta: {result.question}")
    doc.add_heading("Sintesis", level=1)
    for paragraph in result.synthesis.split("\n\n"):
        doc.add_paragraph(paragraph)
    doc.add_heading("Subpreguntas", level=1)
    for item in result.subquestions:
        doc.add_paragraph(item, style="List Bullet")
    if result.formula_candidates:
        doc.add_heading("Formulas / Metodos Detectados", level=1)
        for item in result.formula_candidates:
            doc.add_paragraph(item, style="List Bullet")
    doc.add_heading("Evidencias", level=1)
    for evidence in result.evidences:
        p = doc.add_paragraph()
        p.add_run(f"{evidence.id} - {evidence.title}").bold = True
        doc.add_paragraph(evidence.url)
        doc.add_paragraph(evidence.excerpt)
    doc.add_heading("Lagunas Y Validaciones Pendientes", level=1)
    for item in result.gaps:
        doc.add_paragraph(item, style="List Bullet")
    doc.add_heading("Proximos Pasos", level=1)
    for item in result.next_steps:
        doc.add_paragraph(item, style="List Bullet")
    doc.save(str(output_path))
    return output_path


def export_research_pdf(result: ResearchWorkbenchResult, output_path: Path) -> Path:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    pdf = FPDF(unit="mm", format="A4")
    pdf.set_auto_page_break(auto=True, margin=16)
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 16)
    pdf.cell(0, 10, "ADV ARCHON - Informe de investigacion", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.set_font("Helvetica", "", 9)
    _pdf_multicell(pdf, f"Fecha: {datetime.now().strftime('%Y-%m-%d %H:%M')}")
    _pdf_multicell(pdf, f"Pregunta: {result.question}")
    _pdf_section(pdf, "Sintesis")
    _pdf_multicell(pdf, result.synthesis)
    if result.formula_candidates:
        _pdf_section(pdf, "Formulas / metodos")
        for item in result.formula_candidates:
            _pdf_multicell(pdf, f"- {item}")
    _pdf_section(pdf, "Evidencias")
    for evidence in result.evidences:
        _pdf_multicell(pdf, f"{evidence.id} - {evidence.title}")
        _pdf_multicell(pdf, evidence.url)
        _pdf_multicell(pdf, evidence.excerpt)
    _pdf_section(pdf, "Pendiente de validar")
    for item in result.gaps:
        _pdf_multicell(pdf, f"- {item}")
    pdf.output(str(output_path))
    return output_path


def _default_search(query: str, n: int) -> Any:
    return web_search(query, n=n)


def _default_fetch(url: str) -> Any:
    return web_fetch(url)


def _extract_search_results(raw: Any) -> list[dict[str, Any]]:
    payload = getattr(raw, "payload", raw)
    if not isinstance(payload, dict):
        return []
    results = payload.get("results", [])
    return [item for item in results if isinstance(item, dict)]


def _extract_fetch_text(raw: Any) -> str:
    payload = getattr(raw, "payload", raw)
    if not isinstance(payload, dict):
        return ""
    return str(payload.get("text") or payload.get("content") or "")


def _best_excerpt(text: str, question: str) -> str:
    clean = _clean(text)
    if not clean:
        return ""
    keywords = _keywords(question)
    sentences = re.split(r"(?<=[.!?])\s+", clean)
    scored: list[tuple[int, str]] = []
    for sentence in sentences:
        score = sum(1 for keyword in keywords if keyword in sentence.casefold())
        scored.append((score, sentence))
    scored.sort(key=lambda item: (-item[0], len(item[1])))
    return _clip(scored[0][1] if scored else clean, 700)


def _build_gaps(
    *,
    evidences: Sequence[ResearchEvidence],
    attachments: Sequence[ResearchAttachment],
) -> tuple[str, ...]:
    gaps: list[str] = []
    if not evidences:
        gaps.append("Faltan fuentes web verificadas para sostener la respuesta final.")
    if not attachments:
        gaps.append("No hay documento original adjunto; no se puede contrastar enunciado exacto.")
    gaps.append(
        "Revisar manualmente cualquier formula antes de entregar si tiene impacto "
        "academico o profesional."
    )
    return tuple(gaps)


def _build_next_steps(
    *,
    formula_candidates: Sequence[str],
    evidences: Sequence[ResearchEvidence],
) -> tuple[str, ...]:
    steps = [
        "Redactar respuesta final con estructura de entrega.",
        "Citar fuentes usadas y separar supuestos de resultados.",
    ]
    if formula_candidates:
        steps.append("Pasar las formulas al Spreadsheet Brain para calculo auditable en Excel.")
    if evidences:
        steps.append(
            "Verificar que las fuentes son actuales y apropiadas para el nivel del trabajo."
        )
    return tuple(steps)


def _keywords(text: str) -> list[str]:
    stop = {
        "para", "como", "donde", "cuando", "porque", "sobre", "este", "esta",
        "hacer", "hazme", "dame", "con", "los", "las", "una", "uno", "del",
    }
    words = re.findall(r"[A-Za-zÁÉÍÓÚÜÑáéíóúüñ0-9]{4,}", text.casefold())
    return [word for word in dict.fromkeys(words) if word not in stop][:16]


def _pdf_section(pdf: FPDF, title: str) -> None:
    pdf.ln(3)
    pdf.set_font("Helvetica", "B", 11)
    pdf.cell(0, 7, title, new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.set_font("Helvetica", "", 9)


def _pdf_multicell(pdf: FPDF, text: str) -> None:
    pdf.multi_cell(0, 5, _latin1_safe(text), new_x=XPos.LMARGIN, new_y=YPos.NEXT)


def _clean(text: object) -> str:
    return " ".join(str(text or "").strip().split())


def _clip(text: str, max_chars: int) -> str:
    clean = _clean(text)
    if len(clean) <= max_chars:
        return clean
    return clean[: max(0, max_chars - 15)].rstrip() + " [... recorte]"


def _latin1_safe(text: str) -> str:
    replacements = {
        "—": "-",
        "–": "-",
        "“": '"',
        "”": '"',
        "‘": "'",
        "’": "'",
    }
    for old, new in replacements.items():
        text = text.replace(old, new)
    return text.encode("latin-1", "replace").decode("latin-1")


__all__ = [
    "ResearchAttachment",
    "ResearchEvidence",
    "ResearchWorkbenchResult",
    "build_research_queries",
    "build_research_subquestions",
    "export_research_docx",
    "export_research_pdf",
    "extract_formula_candidates",
    "load_research_attachments",
    "run_research_workbench",
]
