from __future__ import annotations

from collections.abc import Sequence
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook  # type: ignore[import-untyped]
from openpyxl.styles import Font, PatternFill  # type: ignore[import-untyped]
from openpyxl.utils import get_column_letter  # type: ignore[import-untyped]


@dataclass(frozen=True, slots=True)
class SpreadsheetInput:
    name: str
    value: float | int | str
    unit: str = ""
    note: str = ""

    def as_payload(self) -> dict[str, Any]:
        return {
            "name": self.name,
            "value": self.value,
            "unit": self.unit,
            "note": self.note,
        }


@dataclass(frozen=True, slots=True)
class SpreadsheetFormula:
    label: str
    excel_formula: str
    unit: str = ""
    note: str = ""

    def as_payload(self) -> dict[str, str]:
        return {
            "label": self.label,
            "excel_formula": self.excel_formula,
            "unit": self.unit,
            "note": self.note,
        }


@dataclass(frozen=True, slots=True)
class SpreadsheetWorkbookResult:
    output_path: Path
    inputs: tuple[SpreadsheetInput, ...]
    formulas: tuple[SpreadsheetFormula, ...]
    audit_notes: tuple[str, ...]

    def as_payload(self) -> dict[str, Any]:
        return {
            "output_path": str(self.output_path),
            "inputs": [item.as_payload() for item in self.inputs],
            "formulas": [item.as_payload() for item in self.formulas],
            "audit_notes": list(self.audit_notes),
        }


def create_auditable_workbook(
    *,
    title: str,
    inputs: Sequence[SpreadsheetInput],
    formulas: Sequence[SpreadsheetFormula],
    output_path: Path,
    assumptions: Sequence[str] = (),
) -> SpreadsheetWorkbookResult:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws_inputs = wb.active
    ws_inputs.title = "Entradas"
    ws_calc = wb.create_sheet("Calculos")
    ws_assumptions = wb.create_sheet("Supuestos")
    ws_audit = wb.create_sheet("Auditoria")

    _title(ws_inputs, title)
    _headers(ws_inputs, ("Variable", "Valor", "Unidad", "Nota"), row=3)
    for idx, item in enumerate(inputs, start=4):
        ws_inputs.cell(idx, 1, item.name)
        ws_inputs.cell(idx, 2, item.value)
        ws_inputs.cell(idx, 3, item.unit)
        ws_inputs.cell(idx, 4, item.note)

    _title(ws_calc, "Calculos auditables")
    _headers(ws_calc, ("Resultado", "Formula Excel", "Unidad", "Nota"), row=3)
    for idx, formula_item in enumerate(formulas, start=4):
        ws_calc.cell(idx, 1, formula_item.label)
        formula = formula_item.excel_formula
        if formula and not formula.startswith("="):
            formula = "=" + formula
        ws_calc.cell(idx, 2, formula)
        ws_calc.cell(idx, 3, formula_item.unit)
        ws_calc.cell(idx, 4, formula_item.note)

    _title(ws_assumptions, "Supuestos")
    _headers(ws_assumptions, ("#", "Supuesto"), row=3)
    for idx, assumption in enumerate(assumptions or _default_assumptions(), start=4):
        ws_assumptions.cell(idx, 1, idx - 3)
        ws_assumptions.cell(idx, 2, assumption)

    audit_notes = build_workbook_audit_notes(inputs=inputs, formulas=formulas)
    _title(ws_audit, "Auditoria")
    _headers(ws_audit, ("Check", "Resultado"), row=3)
    for idx, note in enumerate(audit_notes, start=4):
        ws_audit.cell(idx, 1, f"A{idx - 3}")
        ws_audit.cell(idx, 2, note)

    for ws in (ws_inputs, ws_calc, ws_assumptions, ws_audit):
        _autosize(ws)
        ws.freeze_panes = "A4"

    wb.save(output_path)
    return SpreadsheetWorkbookResult(
        output_path=output_path,
        inputs=tuple(inputs),
        formulas=tuple(formulas),
        audit_notes=tuple(audit_notes),
    )


def build_workbook_audit_notes(
    *,
    inputs: Sequence[SpreadsheetInput],
    formulas: Sequence[SpreadsheetFormula],
) -> list[str]:
    notes: list[str] = []
    notes.append(f"Entradas declaradas: {len(inputs)}.")
    notes.append(f"Formulas visibles: {len(formulas)}.")
    if not inputs:
        notes.append("Aviso: no hay entradas; el libro puede no ser auditable.")
    if not formulas:
        notes.append("Aviso: no hay formulas; el libro funciona como tabla de datos.")
    for formula in formulas:
        if not formula.excel_formula.strip().startswith("="):
            notes.append(f"Formula '{formula.label}' se normalizara con '=' inicial.")
    notes.append("Revisar unidades y supuestos antes de entregar.")
    return notes


def inspect_workbook_formulas(path: Path) -> tuple[str, ...]:
    wb = load_workbook(path, data_only=False)
    formulas: list[str] = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                if isinstance(value, str) and value.startswith("="):
                    formulas.append(f"{ws.title}!{cell.coordinate}: {value}")
    return tuple(formulas)


def _title(ws: Any, title: str) -> None:
    ws.cell(1, 1, title)
    ws.cell(1, 1).font = Font(bold=True, size=16, color="050505")
    ws.cell(1, 1).fill = PatternFill("solid", fgColor="F7F7F4")


def _headers(ws: Any, labels: Sequence[str], *, row: int) -> None:
    for idx, label in enumerate(labels, start=1):
        cell = ws.cell(row, idx, label)
        cell.font = Font(bold=True, color="050505")
        cell.fill = PatternFill("solid", fgColor="C9A227")


def _autosize(ws: Any) -> None:
    for col_idx, column_cells in enumerate(ws.columns, start=1):
        width = 12
        for cell in column_cells:
            value = str(cell.value or "")
            width = max(width, min(len(value) + 2, 58))
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def _default_assumptions() -> tuple[str, ...]:
    return (
        "Las formulas quedan visibles para auditoria y revision manual.",
        "Las unidades deben revisarse antes de usar el resultado en una entrega formal.",
        "El libro no sustituye criterio profesional ni comprobacion normativa.",
    )


__all__ = [
    "SpreadsheetFormula",
    "SpreadsheetInput",
    "SpreadsheetWorkbookResult",
    "build_workbook_audit_notes",
    "create_auditable_workbook",
    "inspect_workbook_formulas",
]
