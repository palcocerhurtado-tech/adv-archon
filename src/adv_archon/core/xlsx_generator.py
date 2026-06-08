# mypy: ignore-errors
from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any

from adv_archon.core.document_draft import DocumentDraft, DraftTable


def generate_expediente_xlsx(
    draft: DocumentDraft,
    output_path: Path,
) -> Path:
    """Generate an auditable XLSX workbook for an expediente draft."""

    from openpyxl import Workbook  # type: ignore[import-untyped]
    from openpyxl.styles import Font, PatternFill  # type: ignore[import-untyped]
    from openpyxl.utils import get_column_letter  # type: ignore[import-untyped]

    output_path = output_path.expanduser()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Portada"
    checklist = workbook.create_sheet("Checklist")
    sources = workbook.create_sheet("Fuentes")
    risks = workbook.create_sheet("Riesgos")
    tables = workbook.create_sheet("Tablas")
    audit = workbook.create_sheet("Auditoria")

    styles = _Styles(Font=Font, PatternFill=PatternFill)
    _write_summary(summary, draft, styles)
    _write_checklist(checklist, draft, styles)
    _write_sources(sources, draft, styles)
    _write_risks(risks, draft, styles)
    _write_tables(tables, draft, styles)
    _write_audit(audit, draft, styles)

    for sheet in workbook.worksheets:
        _autosize(sheet, get_column_letter=get_column_letter)
        sheet.freeze_panes = "A3"

    workbook.save(output_path)
    return output_path


def generate_research_xlsx(
    draft: DocumentDraft,
    output_path: Path,
) -> Path:
    return generate_expediente_xlsx(draft, output_path)


def generate_budget_xlsx(
    budget: dict[str, Any],
    output_path: Path,
) -> Path:
    from openpyxl import Workbook  # type: ignore[import-untyped]
    from openpyxl.styles import Font, PatternFill  # type: ignore[import-untyped]
    from openpyxl.utils import get_column_letter  # type: ignore[import-untyped]

    output_path = output_path.expanduser()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Presupuesto"
    styles = _Styles(Font=Font, PatternFill=PatternFill)
    _title(sheet, "Presupuesto ADV ARCHON", styles)
    headers = ("Concepto", "Cantidad", "Precio unitario", "Subtotal", "IVA", "Total")
    _headers(sheet, headers, row=3, styles=styles)
    items = budget.get("items", [])
    if not isinstance(items, list):
        items = []
    start_row = 4
    for idx, item in enumerate(items, start=start_row):
        if not isinstance(item, dict):
            continue
        sheet.cell(idx, 1, str(item.get("concept") or item.get("concepto") or ""))
        sheet.cell(idx, 2, float(item.get("quantity") or item.get("cantidad") or 0))
        sheet.cell(idx, 3, float(item.get("unit_price") or item.get("precio_unitario") or 0))
        sheet.cell(idx, 4, f"=B{idx}*C{idx}")
        sheet.cell(idx, 5, f"=D{idx}*0.21")
        sheet.cell(idx, 6, f"=D{idx}+E{idx}")
    total_row = max(start_row, start_row + len(items))
    sheet.cell(total_row, 5, "TOTAL")
    sheet.cell(total_row, 6, f"=SUM(F{start_row}:F{total_row - 1})")
    sheet.cell(total_row, 5).font = styles.header_font
    sheet.cell(total_row, 6).font = styles.header_font
    _autosize(sheet, get_column_letter=get_column_letter)
    workbook.save(output_path)
    return output_path


class _Styles:
    def __init__(self, *, Font: Any, PatternFill: Any) -> None:
        self.title_font = Font(bold=True, size=16, color="C9A227")
        self.header_font = Font(bold=True, color="F7F7F4")
        self.header_fill = PatternFill("solid", fgColor="2E2E2C")
        self.gold_fill = PatternFill("solid", fgColor="F3E7B6")
        self.ok_fill = PatternFill("solid", fgColor="DDEBDD")
        self.review_fill = PatternFill("solid", fgColor="F7E8C9")
        self.blocked_fill = PatternFill("solid", fgColor="F0D4D1")


def _write_summary(sheet: Any, draft: DocumentDraft, styles: _Styles) -> None:
    _title(sheet, "ADV ARCHON - Portada", styles)
    rows = (
        ("Titulo", draft.title),
        ("Cliente", draft.client_name),
        ("Proyecto", draft.project_name),
        ("Municipio", draft.municipality),
        ("Referencia catastral", draft.cadastral_ref),
        ("Veredicto preliminar", draft.verdict),
        ("Fecha", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("Resumen ejecutivo", draft.executive_summary),
    )
    for idx, (key, value) in enumerate(rows, start=3):
        sheet.cell(idx, 1, key)
        sheet.cell(idx, 2, value or "-")
    sheet.cell(8, 2).fill = _status_fill(draft.verdict, styles)


def _write_checklist(sheet: Any, draft: DocumentDraft, styles: _Styles) -> None:
    _headers(sheet, ("Check", "Estado", "Evidencia", "Accion recomendada"), row=1, styles=styles)
    checklist = _find_table(draft, "checklist")
    if checklist:
        for row_idx, row in enumerate(checklist.rows, start=2):
            for col_idx, value in enumerate(row[:4], start=1):
                sheet.cell(row_idx, col_idx, value)
            sheet.cell(row_idx, 2).fill = _status_fill(sheet.cell(row_idx, 2).value, styles)
        return
    default_rows = (
        ("PGOU municipal", "REVISAR", "Pendiente de validacion", "Confirmar ordenanza"),
        ("Catastro", "REVISAR", draft.cadastral_ref or "Sin referencia", "Resolver parcela"),
        ("Fuentes sectoriales", "REVISAR", "Pendiente", "Ejecutar Autopilot"),
    )
    for row_idx, row in enumerate(default_rows, start=2):
        for col_idx, value in enumerate(row, start=1):
            sheet.cell(row_idx, col_idx, value)
        sheet.cell(row_idx, 2).fill = _status_fill(row[1], styles)


def _write_sources(sheet: Any, draft: DocumentDraft, styles: _Styles) -> None:
    _headers(
        sheet,
        ("Fuente", "Tipo", "Referencia/URL", "Resultado", "Confianza", "Oficial"),
        row=1,
        styles=styles,
    )
    for row_idx, source in enumerate(draft.sources, start=2):
        values = (
            source.title,
            source.source_type,
            source.url or source.reference,
            source.result,
            source.confidence,
            "Si" if source.official else "No",
        )
        for col_idx, value in enumerate(values, start=1):
            sheet.cell(row_idx, col_idx, value)


def _write_risks(sheet: Any, draft: DocumentDraft, styles: _Styles) -> None:
    _headers(sheet, ("Riesgo", "Severidad", "Mitigacion"), row=1, styles=styles)
    warnings = draft.warnings or ("Sin advertencias registradas.",)
    for row_idx, warning in enumerate(warnings, start=2):
        sheet.cell(row_idx, 1, warning)
        sheet.cell(row_idx, 2, "Media" if draft.verdict.upper() != "VIABLE" else "Baja")
        sheet.cell(row_idx, 3, "Revision tecnica y trazabilidad de fuentes.")


def _write_tables(sheet: Any, draft: DocumentDraft, styles: _Styles) -> None:
    row = 1
    for table in draft.tables:
        sheet.cell(row, 1, table.title)
        sheet.cell(row, 1).font = styles.title_font
        row += 1
        _headers(sheet, table.columns or ("Dato",), row=row, styles=styles)
        row += 1
        for table_row in table.rows:
            for col_idx, value in enumerate(table_row, start=1):
                sheet.cell(row, col_idx, value)
            row += 1
        row += 2


def _write_audit(sheet: Any, draft: DocumentDraft, styles: _Styles) -> None:
    _headers(sheet, ("Timestamp", "Accion", "Resultado"), row=1, styles=styles)
    rows = (
        (datetime.now().isoformat(timespec="seconds"), "Borrador generado", draft.title),
        (datetime.now().isoformat(timespec="seconds"), "Veredicto preliminar", draft.verdict),
        (
            datetime.now().isoformat(timespec="seconds"),
            "Fuentes incluidas",
            str(len(draft.sources)),
        ),
    )
    for row_idx, row in enumerate(rows, start=2):
        for col_idx, value in enumerate(row, start=1):
            sheet.cell(row_idx, col_idx, value)


def _title(sheet: Any, value: str, styles: _Styles) -> None:
    sheet.cell(1, 1, value)
    sheet.cell(1, 1).font = styles.title_font


def _headers(sheet: Any, values: tuple[str, ...], *, row: int, styles: _Styles) -> None:
    for col_idx, value in enumerate(values, start=1):
        cell = sheet.cell(row, col_idx, value)
        cell.font = styles.header_font
        cell.fill = styles.header_fill
    if values:
        sheet.auto_filter.ref = f"A{row}:{chr(64 + len(values))}{row}"


def _autosize(sheet: Any, *, get_column_letter: Any) -> None:
    for column_cells in sheet.columns:
        max_length = 0
        column = get_column_letter(column_cells[0].column)
        for cell in column_cells:
            max_length = max(max_length, len(str(cell.value or "")))
        sheet.column_dimensions[column].width = min(max(max_length + 2, 12), 48)


def _find_table(draft: DocumentDraft, table_id_part: str) -> DraftTable | None:
    needle = table_id_part.casefold()
    for table in draft.tables:
        if needle in table.id.casefold() or needle in table.title.casefold():
            return table
    return None


def _status_fill(status: object, styles: _Styles) -> Any:
    text = str(status or "").casefold()
    if any(token in text for token in ("ok", "viable", "ready", "apto")):
        return styles.ok_fill
    if any(token in text for token in ("bloque", "incumple", "rechaz", "alto")):
        return styles.blocked_fill
    return styles.review_fill
