from __future__ import annotations

import io
import os
import re
import unicodedata
import warnings
from contextlib import redirect_stderr, redirect_stdout
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import pdfplumber
import pytesseract  # type: ignore[import-untyped]
import trafilatura
from docx import Document
from openpyxl import load_workbook  # type: ignore[import-untyped]
from pdf2image import convert_from_path
from PIL import Image
from pptx import Presentation
from pypdf import PdfReader

IGNORED_DIRS = {".git", "node_modules", ".venv", "__pycache__"}
TEXT_EXTENSIONS = {
    ".txt",
    ".md",
    ".py",
    ".ts",
    ".js",
    ".json",
    ".yaml",
    ".yml",
    ".toml",
    ".csv",
    ".tsv",
    ".rst",
    ".ini",
}
IMAGE_EXTENSIONS = {
    ".png",
    ".jpg",
    ".jpeg",
    ".tif",
    ".tiff",
    ".bmp",
    ".gif",
    ".webp",
}
COMMON_SEARCH_ROOTS = (
    Path.home() / "Desktop",
    Path.home() / "Documents",
    Path.home() / "Downloads",
    Path.home(),
)
WORD_RE = re.compile(r"[a-zA-Z0-9áéíóúüñÁÉÍÓÚÜÑ]+")
PREVIEW_CHAR_LIMIT = 18000
PDF_PREVIEW_PAGE_LIMIT = 12
PDF_TABLE_PREVIEW_PAGE_LIMIT = 4


@dataclass(slots=True)
class ToolResult:
    name: str
    payload: dict[str, Any]


def read_file(
    path: str,
    start_line: int | None = None,
    end_line: int | None = None,
    preview: bool = False,
) -> ToolResult:
    target = Path(path).expanduser().resolve()
    if not target.exists():
        raise FileNotFoundError(f"Path not found: {target}")
    if target.is_dir():
        raise IsADirectoryError(f"Expected file, got directory: {target}")

    text = _read_by_extension(target, preview=preview)

    lines = text.splitlines()
    start_index = max((start_line or 1) - 1, 0)
    end_index = end_line if end_line is not None else len(lines)
    clipped = lines[start_index:end_index]

    return ToolResult(
        name="read_file",
        payload={
            "path": str(target),
            "start_line": start_line or 1,
            "end_line": end_index,
            "preview": preview,
            "content": "\n".join(clipped),
        },
    )


def list_dir(path: str, depth: int = 1) -> ToolResult:
    target = Path(path).expanduser().resolve()
    if not target.exists():
        raise FileNotFoundError(f"Path not found: {target}")
    if not target.is_dir():
        raise NotADirectoryError(f"Expected directory, got file: {target}")

    entries = list(_walk(target, depth=depth))
    return ToolResult(
        name="list_dir",
        payload={
            "path": str(target),
            "depth": depth,
            "entries": entries,
        },
    )


def find_local(
    query: str,
    path: str | None = None,
    folder_hint: str | None = None,
    kind: str = "file",
    max_results: int = 10,
) -> ToolResult:
    roots = _candidate_search_roots(path=path, folder_hint=folder_hint)
    matches = _search_local_entries(
        query=query,
        roots=roots,
        kind=kind,
        max_results=max_results,
    )
    return ToolResult(
        name="find_local",
        payload={
            "query": query,
            "path": path,
            "folder_hint": folder_hint,
            "kind": kind,
            "matches": matches,
        },
    )


def _walk(path: Path, *, depth: int, prefix: str = "") -> list[str]:
    if depth < 0:
        return []

    entries: list[str] = []
    children = sorted(path.iterdir(), key=lambda item: (not item.is_dir(), item.name.lower()))
    for child in children:
        if child.name in IGNORED_DIRS:
            continue
        marker = "/" if child.is_dir() else ""
        entries.append(f"{prefix}{child.name}{marker}")
        if child.is_dir() and depth > 0:
            entries.extend(_walk(child, depth=depth - 1, prefix=f"{prefix}{child.name}/"))
    return entries


def _candidate_search_roots(*, path: str | None, folder_hint: str | None) -> list[Path]:
    if path:
        target = Path(path).expanduser().resolve()
        return [target] if target.exists() else []

    if folder_hint:
        resolved = _resolve_folder_hint(folder_hint)
        if resolved:
            return resolved

    roots: list[Path] = []
    seen: set[Path] = set()
    for candidate in COMMON_SEARCH_ROOTS:
        if candidate.exists() and candidate not in seen:
            roots.append(candidate)
            seen.add(candidate)
    return roots


def _resolve_folder_hint(folder_hint: str) -> list[Path]:
    normalized_hint = _normalize_name(folder_hint)
    tokens = _tokens_for_match(folder_hint)
    matches: list[tuple[int, Path]] = []

    for root in COMMON_SEARCH_ROOTS:
        if not root.exists():
            continue
        try:
            iterator = os.walk(root)
            for current_root, dirnames, _filenames in iterator:
                current_path = Path(current_root)
                dirnames[:] = [
                    name
                    for name in dirnames
                    if name not in IGNORED_DIRS and _depth_from(root, current_path / name) <= 5
                ]
                if current_path == root:
                    continue
                candidate_name = _normalize_name(current_path.name)
                score = _score_name_match(
                    candidate_name,
                    normalized_hint=normalized_hint,
                    tokens=tokens,
                )
                if score > 0:
                    matches.append((score, current_path))
        except OSError:
            continue

    matches.sort(key=lambda item: (-item[0], len(str(item[1]))))
    return [path for _score, path in matches[:5]]


def _search_local_entries(
    *,
    query: str,
    roots: list[Path],
    kind: str,
    max_results: int,
) -> list[dict[str, Any]]:
    normalized_query = _normalize_name(query)
    tokens = _tokens_for_match(query)
    matches: list[tuple[int, Path]] = []

    for root in roots:
        if not root.exists():
            continue
        try:
            for current_root, dirnames, filenames in os.walk(root):
                current_path = Path(current_root)
                dirnames[:] = [
                    name
                    for name in dirnames
                    if name not in IGNORED_DIRS and _depth_from(root, current_path / name) <= 8
                ]
                if kind in {"any", "dir"}:
                    for dirname in dirnames:
                        candidate = current_path / dirname
                        score = _score_name_match(
                            _normalize_name(dirname),
                            normalized_hint=normalized_query,
                            tokens=tokens,
                        )
                        if score > 0:
                            matches.append((score, candidate))
                if kind in {"any", "file"}:
                    for filename in filenames:
                        candidate = current_path / filename
                        score = _score_name_match(
                            _normalize_name(candidate.stem),
                            normalized_hint=normalized_query,
                            tokens=tokens,
                        )
                        if score > 0:
                            matches.append((score, candidate))
        except OSError:
            continue

    matches.sort(key=lambda item: (-item[0], len(str(item[1]))))
    seen: set[Path] = set()
    results: list[dict[str, Any]] = []
    for score, match in matches:
        if match in seen:
            continue
        seen.add(match)
        results.append(
            {
                "path": str(match),
                "name": match.name,
                "is_dir": match.is_dir(),
                "score": score,
            }
        )
        if len(results) >= max_results:
            break
    return results


def _score_name_match(candidate_name: str, *, normalized_hint: str, tokens: list[str]) -> int:
    score = 0
    if normalized_hint and normalized_hint in candidate_name:
        score += 10
    if tokens and all(token in candidate_name for token in tokens):
        score += 8
    score += sum(1 for token in tokens if token in candidate_name)
    return score


def _tokens_for_match(text: str) -> list[str]:
    return [
        _normalize_name(match.group(0))
        for match in WORD_RE.finditer(text)
        if len(match.group(0)) > 2
    ]


def _normalize_name(text: str) -> str:
    lowered = unicodedata.normalize("NFKD", text.casefold())
    return "".join(ch for ch in lowered if not unicodedata.combining(ch))


def _depth_from(root: Path, candidate: Path) -> int:
    try:
        return len(candidate.relative_to(root).parts)
    except ValueError:
        return 99


def _read_by_extension(path: Path, *, preview: bool = False) -> str:
    suffix = path.suffix.lower()
    if suffix in TEXT_EXTENSIONS:
        text = path.read_text(encoding="utf-8", errors="replace")
        return _clip_preview_text(text) if preview else text
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        with redirect_stdout(io.StringIO()), redirect_stderr(io.StringIO()):
            if suffix == ".pdf":
                return _read_pdf(path, preview=preview)
            if suffix == ".docx":
                text = _read_docx(path)
                return _clip_preview_text(text) if preview else text
            if suffix == ".xlsx":
                text = _read_xlsx(path)
                return _clip_preview_text(text) if preview else text
            if suffix == ".pptx":
                text = _read_pptx(path)
                return _clip_preview_text(text) if preview else text
            if suffix in {".html", ".htm"}:
                text = _read_html(path)
                return _clip_preview_text(text) if preview else text
            if suffix in IMAGE_EXTENSIONS:
                text = _read_image_with_ocr(path)
                return _clip_preview_text(text) if preview else text
    text = path.read_text(encoding="utf-8", errors="replace")
    return _clip_preview_text(text) if preview else text


def _read_pdf(path: Path, *, preview: bool = False) -> str:
    reader = PdfReader(str(path))
    fragments: list[str] = []
    page_limit = PDF_PREVIEW_PAGE_LIMIT if preview else None
    char_target = PREVIEW_CHAR_LIMIT if preview else None
    total_chars = 0
    for page_index, page in enumerate(reader.pages, start=1):
        if page_limit is not None and page_index > page_limit:
            break
        extracted = page.extract_text() or ""
        if extracted.strip():
            fragments.append(extracted)
            total_chars += len(extracted)
            if char_target is not None and total_chars >= char_target:
                break

    table_text = _read_pdf_tables(
        path,
        page_limit=PDF_TABLE_PREVIEW_PAGE_LIMIT if preview else None,
    )
    if table_text:
        fragments.append(table_text)

    combined = "\n\n".join(fragment for fragment in fragments if fragment.strip())
    if combined.strip():
        return _clip_preview_text(combined) if preview else combined

    ocr_text = _ocr_pdf(path, preview=preview)
    if ocr_text:
        return ocr_text

    return (
        "No he podido extraer texto del PDF. "
        "Si es un PDF escaneado, instala `tesseract` y `poppler` para activar OCR."
    )


def _read_pdf_tables(path: Path, *, page_limit: int | None = None) -> str:
    fragments: list[str] = []
    try:
        with pdfplumber.open(path) as pdf:
            for page_index, page in enumerate(pdf.pages, start=1):
                if page_limit is not None and page_index > page_limit:
                    break
                tables = page.extract_tables()
                for table_index, table in enumerate(tables, start=1):
                    fragments.append(f"Page {page_index}, table {table_index}:")
                    for row in table:
                        values = [
                            cell.strip() if isinstance(cell, str) else str(cell or "")
                            for cell in row
                        ]
                        if any(value for value in values):
                            fragments.append("\t".join(values))
    except Exception:
        return ""
    return "\n".join(fragments)


def _ocr_pdf(path: Path, *, preview: bool = False) -> str:
    try:
        images = convert_from_path(str(path))
    except Exception:
        return ""
    fragments: list[str] = []
    page_limit = PDF_PREVIEW_PAGE_LIMIT if preview else None
    total_chars = 0
    for page_index, image in enumerate(images, start=1):
        if page_limit is not None and page_index > page_limit:
            break
        text = pytesseract.image_to_string(image)
        if text.strip():
            fragments.append(text)
            total_chars += len(text)
            if preview and total_chars >= PREVIEW_CHAR_LIMIT:
                break
    combined = "\n\n".join(fragments)
    return _clip_preview_text(combined) if preview else combined


def _read_docx(path: Path) -> str:
    document = Document(str(path))
    fragments: list[str] = []
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            fragments.append(paragraph.text)
    for table in document.tables:
        for row in table.rows:
            values = [cell.text.strip() for cell in row.cells]
            if any(values):
                fragments.append("\t".join(values))
    return "\n".join(fragments)


def _read_xlsx(path: Path) -> str:
    workbook = load_workbook(path, read_only=True, data_only=True)
    fragments: list[str] = []
    for sheet in workbook.worksheets:
        fragments.append(f"# Sheet: {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            values = ["" if value is None else str(value) for value in row]
            if any(values):
                fragments.append("\t".join(values))
    return "\n".join(fragments)


def _read_pptx(path: Path) -> str:
    presentation = Presentation(str(path))
    fragments: list[str] = []
    for index, slide in enumerate(presentation.slides, start=1):
        fragments.append(f"# Slide {index}")
        for shape in slide.shapes:
            if hasattr(shape, "text") and isinstance(shape.text, str) and shape.text.strip():
                fragments.append(shape.text.strip())
    return "\n".join(fragments)


def _read_html(path: Path) -> str:
    raw = path.read_text(encoding="utf-8", errors="replace")
    extracted = trafilatura.extract(raw, include_links=True, include_tables=True)
    return extracted if extracted else raw


def _read_image_with_ocr(path: Path) -> str:
    with Image.open(path) as image:
        extracted = str(pytesseract.image_to_string(image))
    if extracted.strip():
        return extracted
    return "No he podido extraer texto mediante OCR."


def _clip_preview_text(text: str, *, max_chars: int = PREVIEW_CHAR_LIMIT) -> str:
    compact = text.strip()
    if len(compact) <= max_chars:
        return compact
    head_size = int(max_chars * 0.7)
    tail_size = max_chars - head_size
    head = compact[:head_size].rstrip()
    tail = compact[-tail_size:].lstrip()
    return (
        f"{head}\n\n[... contenido intermedio omitido para agilizar la lectura ...]\n\n{tail}"
    )
