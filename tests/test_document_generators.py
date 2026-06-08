from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from adv_archon.core.document_draft import (
    DocumentDraft,
    DraftSection,
    DraftSource,
    DraftTable,
)
from adv_archon.core.docx_generator import generate_expediente_docx
from adv_archon.core.pdf_generator_v2 import generate_draft_pdf
from adv_archon.core.xlsx_generator import generate_budget_xlsx, generate_expediente_xlsx


def _draft() -> DocumentDraft:
    return DocumentDraft(
        title="Cambio de uso local a vivienda",
        client_name="Estudio Prueba",
        municipality="Madrid",
        cadastral_ref="28079...",
        verdict="CONDICIONADO",
        executive_summary="Viable como primera aproximacion.",
        sections=(DraftSection("summary", "Resumen", "Texto del informe."),),
        tables=(
            DraftTable(
                "checklist",
                "Checklist",
                ("Check", "Estado", "Evidencia", "Accion"),
                (("Catastro", "OK", "Ref verificada", "Mantener"),),
            ),
        ),
        sources=(
            DraftSource(
                "catastro",
                "Catastro OVC",
                "catastro",
                result="Referencia verificada",
                official=True,
            ),
        ),
        warnings=("Revisar ordenanza.",),
        next_steps=("Confirmar PGOU.",),
    )


def test_generate_expediente_docx_creates_word_file(tmp_path: Path) -> None:
    output = generate_expediente_docx(_draft(), tmp_path / "informe.docx")

    assert output.exists()
    assert output.stat().st_size > 1000


def test_generate_expediente_xlsx_creates_auditable_workbook(tmp_path: Path) -> None:
    output = generate_expediente_xlsx(_draft(), tmp_path / "informe.xlsx")

    workbook = load_workbook(output, data_only=False)

    assert "Portada" in workbook.sheetnames
    assert "Fuentes" in workbook.sheetnames
    assert workbook["Portada"]["B3"].value == "Cambio de uso local a vivienda"
    assert workbook["Checklist"]["A2"].value == "Catastro"


def test_generate_budget_xlsx_keeps_formulas(tmp_path: Path) -> None:
    output = generate_budget_xlsx(
        {
            "items": [
                {"concept": "Informe", "quantity": 2, "unit_price": 100},
            ]
        },
        tmp_path / "presupuesto.xlsx",
    )

    workbook = load_workbook(output, data_only=False)

    assert workbook["Presupuesto"]["D4"].value == "=B4*C4"
    assert workbook["Presupuesto"]["F4"].value == "=D4+E4"


def test_generate_draft_pdf_creates_pdf(tmp_path: Path) -> None:
    output = generate_draft_pdf(_draft(), tmp_path / "informe.pdf")

    assert output.exists()
    assert output.stat().st_size > 1000
