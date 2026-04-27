from pathlib import Path
from unittest.mock import patch

import pytest
from docx import Document
from openpyxl import Workbook
from pptx import Presentation

from adv_archon.tools.files import find_local, list_dir, read_file


def test_read_file_text(tmp_path: Path) -> None:
    target = tmp_path / "note.txt"
    target.write_text("one\ntwo\nthree\n", encoding="utf-8")

    result = read_file(str(target), start_line=2, end_line=3)

    assert result.payload["content"] == "two\nthree"


def test_list_dir_ignores_noise(tmp_path: Path) -> None:
    (tmp_path / ".git").mkdir()
    (tmp_path / "docs").mkdir()
    (tmp_path / "docs" / "readme.md").write_text("hello", encoding="utf-8")

    result = list_dir(str(tmp_path), depth=2)

    assert ".git/" not in result.payload["entries"]
    assert "docs/" in result.payload["entries"]
    assert "docs/readme.md" in result.payload["entries"]


def test_read_file_docx(tmp_path: Path) -> None:
    target = tmp_path / "brief.docx"
    document = Document()
    document.add_paragraph("Linea uno")
    document.add_paragraph("Linea dos")
    document.save(target)

    result = read_file(str(target))

    assert "Linea uno" in result.payload["content"]
    assert "Linea dos" in result.payload["content"]


def test_read_file_xlsx(tmp_path: Path) -> None:
    target = tmp_path / "sheet.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Plan"
    sheet["A1"] = "Tarea"
    sheet["B1"] = "Estado"
    sheet["A2"] = "Bloque 3"
    sheet["B2"] = "activo"
    workbook.save(target)

    result = read_file(str(target))

    assert "# Sheet: Plan" in result.payload["content"]
    assert "Bloque 3\tactivo" in result.payload["content"]


def test_read_file_pptx(tmp_path: Path) -> None:
    target = tmp_path / "deck.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Demo"
    slide.placeholders[1].text = "Contenido"
    presentation.save(target)

    result = read_file(str(target))

    assert "# Slide 1" in result.payload["content"]
    assert "Demo" in result.payload["content"]


def test_read_file_suppresses_noisy_document_output(
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
) -> None:
    target = tmp_path / "brief.pdf"
    target.write_text("fake", encoding="utf-8")

    def fake_read_pdf_ok(_path: Path, *, preview: bool = False) -> str:
        import sys

        sys.stdout.write("ruido stdout")
        sys.stderr.write("ruido stderr")
        return "contenido pdf"

    with patch("adv_archon.tools.files._read_pdf", side_effect=fake_read_pdf_ok):
        result = read_file(str(target))

    captured = capsys.readouterr()

    assert result.payload["content"] == "contenido pdf"
    assert captured.out == ""
    assert captured.err == ""


def test_read_file_preview_clips_large_text(tmp_path: Path) -> None:
    target = tmp_path / "large.txt"
    target.write_text(("intro " * 5000) + ("cierre " * 2000), encoding="utf-8")

    result = read_file(str(target), preview=True)

    assert result.payload["preview"] is True
    assert "[... contenido intermedio omitido para agilizar la lectura ...]" in result.payload[
        "content"
    ]


def test_find_local_resolves_folder_hint_and_book_title(tmp_path: Path) -> None:
    desktop = tmp_path / "Desktop"
    desktop.mkdir()
    grasa = desktop / "LA GRASA"
    grasa.mkdir()
    target = grasa / "La Clavicula de Salomon.pdf"
    target.write_text("contenido", encoding="utf-8")

    with patch(
        "adv_archon.tools.files.COMMON_SEARCH_ROOTS",
        (desktop, tmp_path / "Documents", tmp_path / "Downloads", tmp_path),
    ):
        result = find_local(
            query="la clavicula de salomon",
            folder_hint="la grasa",
            kind="file",
            max_results=5,
        )

    matches = result.payload["matches"]
    assert matches
    assert matches[0]["path"] == str(target)
    assert matches[0]["is_dir"] is False
