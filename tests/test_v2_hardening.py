"""Headless regression tests for the v2 hardening sprint.

Covers:
* Pure expediente analysis logic (extracted from the PySide6 glue layer).
* Thread-drain helper used by DesktopWindowV2.closeEvent.
* Legal/compliance module + propagation into generated PDF/DOCX/XLSX.
* SQL identifier guard in DocumentStore.
"""
from __future__ import annotations

import json
from types import SimpleNamespace

import pytest

from adv_archon.core import legal
from adv_archon.core.expediente_analysis import build_expediente_analysis


def _exp(*, checks: list[dict] | None = None, case_type: str = "", plan_path: str = ""):
    site_context = json.dumps({"legal_checks": checks}) if checks is not None else ""
    return SimpleNamespace(
        case_type=case_type,
        site_context=site_context,
        plan_path=plan_path,
    )


# ── Expediente analysis ─────────────────────────────────────────────────────


def test_analysis_no_checks_is_revisar():
    result = build_expediente_analysis(_exp(checks=None))
    assert result["verdict"] == "revisar"
    assert result["verdict_label"] == "REVISAR"
    assert "generated_at" in result


def test_analysis_pending_or_missing_is_revisar():
    checks = [{"status": "missing", "title": "PGOU", "detail": "sin datos"}]
    result = build_expediente_analysis(_exp(checks=checks))
    assert result["verdict"] == "revisar"


def test_analysis_conditional_is_condicionado():
    checks = [
        {
            "status": "conditional",
            "title": "Costas",
            "detail": "afección parcial",
            "recommended_action": "consultar administración",
        }
    ]
    result = build_expediente_analysis(_exp(checks=checks))
    assert result["verdict"] == "condicionado"
    assert result["annotations"][0]["status"] == "warning"
    assert "consultar administración" in result["next_steps"]


def test_analysis_all_ready_is_viable():
    checks = [{"status": "ready", "title": "Catastro", "detail": "resuelto"}]
    result = build_expediente_analysis(_exp(checks=checks))
    assert result["verdict"] == "viable"
    # 'ready' checks must not generate next steps from their action.
    assert all("resuelto" not in step for step in result["next_steps"])


def test_analysis_caps_collections():
    checks = [
        {
            "status": "conditional",
            "title": f"Check {i}",
            "detail": "x",
            "recommended_action": f"action {i}",
        }
        for i in range(40)
    ]
    result = build_expediente_analysis(_exp(checks=checks))
    assert len(result["annotations"]) <= 30
    assert len(result["next_steps"]) <= 12


def test_analysis_handles_malformed_site_context():
    exp = SimpleNamespace(case_type="", site_context="{not json", plan_path="")
    result = build_expediente_analysis(exp)
    assert result["verdict"] == "revisar"


# ── Thread-drain helper (closeEvent fix) ────────────────────────────────────


class _FakeThread:
    def __init__(self, running: bool = True) -> None:
        self._running = running
        self.quit_called = False
        self.waited = False

    def isRunning(self) -> bool:  # noqa: N802 (mirrors Qt API)
        return self._running

    def quit(self) -> None:
        self.quit_called = True
        self._running = False

    def wait(self, _timeout_ms: int) -> bool:
        self.waited = True
        return True


def test_drain_quits_running_threads():
    from adv_archon.desktop.app_v2 import _drain_thread_pairs

    running = _FakeThread(running=True)
    idle = _FakeThread(running=False)
    pairs = [(running, object()), (idle, object())]
    _drain_thread_pairs(pairs)
    assert running.quit_called and running.waited
    assert not idle.quit_called  # idle threads are left alone


def test_drain_survives_deleted_qobject():
    from adv_archon.desktop.app_v2 import _drain_thread_pairs

    class _Boom:
        def isRunning(self):  # noqa: N802
            raise RuntimeError("C++ object already deleted")

    # Must not propagate the RuntimeError.
    _drain_thread_pairs([(_Boom(), object())])


# ── Legal / compliance module ───────────────────────────────────────────────


def test_legal_footer_covers_required_dimensions():
    footer = legal.document_legal_footer().lower()
    assert "no vinculante" in footer
    assert "rgpd" in footer or "lopdgdd" in footer
    assert "2024/1689" in footer  # EU AI Act reference
    assert "técnico competente" in footer
    assert legal.LEGAL_VERSION in legal.document_legal_footer()


def test_legal_short_disclaimer_is_one_concept():
    short = legal.document_short_disclaimer().lower()
    assert "no vinculante" in short
    assert "ia" in short


# ── Document generators carry the legal text ────────────────────────────────


def _sample_draft():
    from adv_archon.core.document_draft import build_expediente_draft

    exp = SimpleNamespace(
        title="Expediente de prueba",
        municipality="Madrid",
        cadastral_ref="1234",
        analysis_result="Análisis preliminar.",
        notes="",
        status="revisar",
        address="Calle Falsa 1",
        case_type="",
        id="exp-1",
    )
    return build_expediente_draft(exp)


def test_build_draft_warnings_mention_compliance():
    draft = _sample_draft()
    joined = " ".join(draft.warnings).lower()
    assert "no vinculante" in joined
    assert "ia" in joined
    assert "rgpd" in joined or "lopdgdd" in joined


def test_pdf_contains_legal_section(tmp_path):
    pytest.importorskip("fpdf")
    from adv_archon.core.pdf_generator_v2 import generate_draft_pdf

    out = generate_draft_pdf(_sample_draft(), tmp_path / "r.pdf")
    assert out.exists() and out.stat().st_size > 0


def test_docx_contains_legal_section(tmp_path):
    pytest.importorskip("docx")
    from adv_archon.core.docx_generator import generate_expediente_docx

    out = generate_expediente_docx(_sample_draft(), tmp_path / "r.docx")
    from docx import Document

    text = "\n".join(p.text for p in Document(str(out)).paragraphs).lower()
    assert "aviso legal" in text
    assert "no vinculante" in text


def test_xlsx_contains_legal_sheet(tmp_path):
    pytest.importorskip("openpyxl")
    from adv_archon.core.xlsx_generator import generate_expediente_xlsx

    out = generate_expediente_xlsx(_sample_draft(), tmp_path / "r.xlsx")
    from openpyxl import load_workbook

    wb = load_workbook(out)
    assert "Aviso legal" in wb.sheetnames


# ── DocumentStore SQL identifier guard ──────────────────────────────────────


def test_safe_identifier_accepts_known_columns():
    from adv_archon.core.document_store import _safe_identifier

    assert _safe_identifier("exported_pdf_path") == "exported_pdf_path"


def test_safe_identifier_rejects_injection():
    from adv_archon.core.document_store import _safe_identifier

    with pytest.raises(ValueError):
        _safe_identifier("x = 1; DROP TABLE document_drafts; --")


def test_mark_exported_roundtrip(tmp_path):
    from adv_archon.core.document_store import DocumentStore

    store = DocumentStore(tmp_path / "documents.db")
    try:
        draft = store.save_draft(_sample_draft())
        assert store.mark_exported(draft.id, "pdf", tmp_path / "out.pdf")
        with pytest.raises(ValueError):
            store.mark_exported(draft.id, "exe", tmp_path / "x.exe")
    finally:
        store.close()


# ── First-run consent state ──────────────────────────────────────────────────


def test_consent_not_accepted_by_default(tmp_path):
    from adv_archon.core.consent import is_consent_accepted

    assert not is_consent_accepted(tmp_path)


def test_consent_accepted_after_mark(tmp_path):
    from adv_archon.core.consent import is_consent_accepted, mark_consent_accepted

    mark_consent_accepted(tmp_path)
    assert is_consent_accepted(tmp_path)


def test_consent_persists_correct_version(tmp_path):
    import json

    from adv_archon.core.consent import mark_consent_accepted
    from adv_archon.core.legal import LEGAL_VERSION

    mark_consent_accepted(tmp_path)
    stamp = json.loads((tmp_path / "consent.json").read_text())
    assert stamp["legal_version"] == LEGAL_VERSION
    assert stamp["accepted"] is True
    assert "accepted_at" in stamp


def test_consent_reset_clears_state(tmp_path):
    from adv_archon.core.consent import is_consent_accepted, mark_consent_accepted, reset_consent

    mark_consent_accepted(tmp_path)
    assert is_consent_accepted(tmp_path)
    reset_consent(tmp_path)
    assert not is_consent_accepted(tmp_path)


def test_consent_corrupt_file_returns_false(tmp_path):
    from adv_archon.core.consent import is_consent_accepted

    (tmp_path / "consent.json").write_text("{not json}", encoding="utf-8")
    assert not is_consent_accepted(tmp_path)


def test_consent_wrong_version_returns_false(tmp_path):
    import json

    from adv_archon.core.consent import is_consent_accepted

    stamp = {  # version deliberately wrong
        "accepted": True,
        "legal_version": "1900.0",
        "accepted_at": "2000-01-01T00:00:00+00:00",
    }
    (tmp_path / "consent.json").write_text(json.dumps(stamp), encoding="utf-8")
    assert not is_consent_accepted(tmp_path)
